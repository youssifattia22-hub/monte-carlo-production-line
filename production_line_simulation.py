
"""
Production-line simulation V2
=============================
Reads fitted distribution parameters directly from Excel "Model Inputs",
generates the random values in Python, runs the M1 -> B1 -> M2 -> B2 -> M3
logic, and includes deterministic validation tests.

Excel expected layout (current Project.xlsx):
TTF:   rows 4:6   -> adjusted mu I, sigma G
TTR:   rows 9:11  -> shape F, adjusted scale I
PT:    rows 14:16 -> adjusted min I, adjusted max J
Scrap: rows 19:21 -> adjusted mu I, sigma G

Requires:
    pip install openpyxl
"""

from dataclasses import dataclass
from pathlib import Path
from typing import Optional
import argparse, csv, math, random, statistics, os, time
from concurrent.futures import ProcessPoolExecutor

try:
    from openpyxl import load_workbook
except ImportError:
    raise SystemExit("Missing openpyxl. Install it with: pip install openpyxl")


@dataclass
class MachineParams:
    name: str
    ttf_mu: float
    ttf_sigma: float
    ttr_shape: float
    ttr_scale: float
    pt_min_sec: float
    pt_max_sec: float
    scrap_mu: float
    scrap_sigma: float


@dataclass
class Scenario:
    buffer1_capacity: int = 10
    buffer2_capacity: int = 10
    pm_interval_h: Optional[float] = 200.0
    pm_duration_h: float = 1.0
    shift_duration_h: float = 8.0
    demand_per_shift: float = 300.0
    horizon_h: float = 720.0


@dataclass
class MachineState:
    p: MachineParams
    state: str = "IDLE"
    remaining_h: float = 0.0
    operating_age_h: float = 0.0
    next_failure_age_h: float = math.inf
    completed: int = 0
    good_completed: int = 0
    scrap: int = 0
    failures: int = 0
    pm_events: int = 0
    repair_h: float = 0.0
    pm_h: float = 0.0
    processing_h: float = 0.0
    blocked_h: float = 0.0
    starved_h: float = 0.0
    idle_h: float = 0.0


def _num(v, label):
    if v is None or not isinstance(v, (int, float)):
        raise ValueError(
            f"{label} is not a numeric cached Excel value ({v!r}). "
            "Open Project.xlsx in Excel, let formulas recalculate, Save, then rerun."
        )
    return float(v)


def load_inputs_from_excel(path: str) -> list[MachineParams]:
    """Read the CURRENT calculated values from Model Inputs."""
    wb = load_workbook(path, data_only=True, read_only=True)
    if "Model Inputs" not in wb.sheetnames:
        raise ValueError("Sheet 'Model Inputs' was not found.")
    ws = wb["Model Inputs"]

    machines = []
    for i, (r_ttf, r_ttr, r_pt, r_scrap) in enumerate(
        zip((4,5,6), (9,10,11), (14,15,16), (19,20,21)), start=1
    ):
        machines.append(MachineParams(
            name=f"M{i}",
            ttf_mu=_num(ws[f"I{r_ttf}"].value, f"M{i} TTF adjusted mu"),
            ttf_sigma=_num(ws[f"G{r_ttf}"].value, f"M{i} TTF sigma"),
            ttr_shape=_num(ws[f"F{r_ttr}"].value, f"M{i} TTR shape"),
            ttr_scale=_num(ws[f"I{r_ttr}"].value, f"M{i} TTR adjusted scale"),
            pt_min_sec=_num(ws[f"I{r_pt}"].value, f"M{i} PT adjusted min"),
            pt_max_sec=_num(ws[f"J{r_pt}"].value, f"M{i} PT adjusted max"),
            scrap_mu=_num(ws[f"I{r_scrap}"].value, f"M{i} scrap adjusted mu"),
            scrap_sigma=_num(ws[f"G{r_scrap}"].value, f"M{i} scrap sigma"),
        ))
    return machines


def sample_ttf(p, rng):
    return rng.lognormvariate(p.ttf_mu, p.ttf_sigma)

def sample_ttr(p, rng):
    return rng.weibullvariate(p.ttr_scale, p.ttr_shape)

def sample_pt_h(p, rng):
    return rng.uniform(p.pt_min_sec, p.pt_max_sec) / 3600.0

def sample_scrap_probability(p, rng):
    return min(1.0, max(0.0, rng.lognormvariate(p.scrap_mu, p.scrap_sigma)))

def init_machine(p, rng):
    m = MachineState(p=p)
    m.next_failure_age_h = sample_ttf(p, rng)
    return m

def start_processing(m, rng):
    m.state = "PROCESSING"
    m.remaining_h = sample_pt_h(m.p, rng)

def reset_after_maintenance(m, rng):
    m.operating_age_h = 0.0
    m.next_failure_age_h = sample_ttf(m.p, rng)
    m.state = "IDLE"
    m.remaining_h = 0.0


def run_one(params, s: Scenario, seed: int):
    rng = random.Random(seed)
    m1, m2, m3 = [init_machine(p, rng) for p in params]
    machines = [m1, m2, m3]
    b1 = b2 = finished_good = 0
    # Counts accepted good transfers for conservation checks.
    m1_to_b1 = m2_to_b2 = 0
    t = 0.0
    eps = 1e-10

    while t < s.horizon_h - eps:
        # Downstream-first starts free buffer capacity before upstream decisions.
        if m3.state == "IDLE" and b2 > 0:
            b2 -= 1
            start_processing(m3, rng)

        if m2.state == "IDLE" and b1 > 0 and b2 < s.buffer2_capacity:
            b1 -= 1
            start_processing(m2, rng)

        if m1.state == "IDLE" and b1 < s.buffer1_capacity:
            start_processing(m1, rng)

        candidates = [s.horizon_h - t]
        for m in machines:
            if m.state in ("PROCESSING", "REPAIR", "PM"):
                candidates.append(max(eps, m.remaining_h))
            if m.state == "PROCESSING":
                candidates.append(max(eps, m.next_failure_age_h - m.operating_age_h))
                if s.pm_interval_h is not None:
                    candidates.append(max(eps, s.pm_interval_h - m.operating_age_h))
        dt = min(candidates)

        for idx, m in enumerate(machines):
            if m.state == "PROCESSING":
                m.processing_h += dt
                m.operating_age_h += dt
                m.remaining_h -= dt
            elif m.state == "REPAIR":
                m.repair_h += dt
                m.remaining_h -= dt
            elif m.state == "PM":
                m.pm_h += dt
                m.remaining_h -= dt
            elif m.state == "IDLE":
                if idx == 0:
                    m.blocked_h += dt if b1 >= s.buffer1_capacity else 0
                    m.idle_h += dt if b1 < s.buffer1_capacity else 0
                elif idx == 1:
                    if b1 <= 0: m.starved_h += dt
                    elif b2 >= s.buffer2_capacity: m.blocked_h += dt
                    else: m.idle_h += dt
                else:
                    if b2 <= 0: m.starved_h += dt
                    else: m.idle_h += dt
        t += dt

        # Failure/PM priority. Assumption: interrupted unit is lost as process loss.
        # This makes the handling explicit rather than silently resuming it.
        for m in machines:
            if m.state == "PROCESSING" and m.operating_age_h >= m.next_failure_age_h - eps:
                m.state = "REPAIR"
                m.remaining_h = sample_ttr(m.p, rng)
                m.failures += 1
                continue
            if (m.state == "PROCESSING" and s.pm_interval_h is not None
                    and m.operating_age_h >= s.pm_interval_h - eps):
                m.state = "PM"
                m.remaining_h = s.pm_duration_h
                m.pm_events += 1

        for m in machines:
            if m.state == "REPAIR" and m.remaining_h <= eps:
                reset_after_maintenance(m, rng)
            elif m.state == "PM" and m.remaining_h <= eps:
                reset_after_maintenance(m, rng)

        # Completed processing -> decide scrap after processing.
        if m3.state == "PROCESSING" and m3.remaining_h <= eps:
            m3.completed += 1
            if rng.random() < sample_scrap_probability(m3.p, rng):
                m3.scrap += 1
            else:
                m3.good_completed += 1
                finished_good += 1
            m3.state = "IDLE"

        if m2.state == "PROCESSING" and m2.remaining_h <= eps:
            m2.completed += 1
            if rng.random() < sample_scrap_probability(m2.p, rng):
                m2.scrap += 1
            else:
                m2.good_completed += 1
                if b2 >= s.buffer2_capacity:
                    raise RuntimeError("B2 overflow: logic error.")
                b2 += 1
                m2_to_b2 += 1
            m2.state = "IDLE"

        if m1.state == "PROCESSING" and m1.remaining_h <= eps:
            m1.completed += 1
            if rng.random() < sample_scrap_probability(m1.p, rng):
                m1.scrap += 1
            else:
                m1.good_completed += 1
                if b1 >= s.buffer1_capacity:
                    raise RuntimeError("B1 overflow: logic error.")
                b1 += 1
                m1_to_b1 += 1
            m1.state = "IDLE"

        if not (0 <= b1 <= s.buffer1_capacity and 0 <= b2 <= s.buffer2_capacity):
            raise RuntimeError("Buffer bounds violated.")

    demand_target = s.demand_per_shift * (s.horizon_h / s.shift_duration_h)
    out = {
        "seed": seed, "finished_good": finished_good,
        "demand_target": demand_target,
        "demand_met": int(finished_good >= demand_target),
        "throughput_good_per_h": finished_good / s.horizon_h,
        "ending_b1": b1, "ending_b2": b2,
        "M1_to_B1": m1_to_b1, "M2_to_B2": m2_to_b2,
    }
    for m in machines:
        for key in ("completed","good_completed","scrap","failures","pm_events",
                    "repair_h","pm_h","processing_h","blocked_h","starved_h"):
            out[f"{m.p.name}_{key}"] = getattr(m, key)
        out[f"{m.p.name}_availability"] = 1 - (m.repair_h + m.pm_h) / s.horizon_h
    return out


def _run_worker(task):
    params, scenario, seed = task
    return run_one(params, scenario, seed)

# ---------- deterministic verification tests ----------

def deterministic_params(pt_sec=(60,60,60), scrap_mu=-100.0):
    # Huge TTF => no failures; almost-zero scrap.
    return [
        MachineParams(f"M{i+1}", math.log(1e9), 0.0, 2.0, 1.0,
                      pt_sec[i], pt_sec[i], scrap_mu, 0.0)
        for i in range(3)
    ]

def validation_suite():
    tests = []

    # Test 1: no failures, no scrap, equal 60-sec machines.
    # In one hour, pipeline latency means 58 finished units (first M3 completion at t=3 min).
    p = deterministic_params((60,60,60))
    r = run_one(p, Scenario(20,20,None,horizon_h=1,demand_per_shift=0), 1)
    tests.append(("Deterministic flow", r["finished_good"] == 58,
                  f"finished_good={r['finished_good']} expected=58"))

    # Test 2: upstream much faster -> M1 must be blocked.
    p = deterministic_params((10,120,120))
    r = run_one(p, Scenario(2,20,None,horizon_h=1,demand_per_shift=0), 2)
    tests.append(("Blocking", r["M1_blocked_h"] > 0,
                  f"M1_blocked_h={r['M1_blocked_h']:.4f}"))

    # Test 3: M1 slow, M2/M3 fast -> starvation must occur.
    p = deterministic_params((120,10,10))
    r = run_one(p, Scenario(20,20,None,horizon_h=1,demand_per_shift=0), 3)
    tests.append(("Starvation", r["M2_starved_h"] > 0 and r["M3_starved_h"] > 0,
                  f"M2={r['M2_starved_h']:.4f}, M3={r['M3_starved_h']:.4f}"))

    # Test 4: buffer bounds are always enforced internally; reaching here is a pass.
    tests.append(("Buffer bounds", 0 <= r["ending_b1"] <= 20 and 0 <= r["ending_b2"] <= 20,
                  f"B1={r['ending_b1']}, B2={r['ending_b2']}"))

    # Test 5: completed pieces split exactly into good + scrap.
    ok = all(r[f"M{i}_completed"] == r[f"M{i}_good_completed"] + r[f"M{i}_scrap"]
             for i in (1,2,3))
    tests.append(("Completion conservation", ok, "completed = good + scrap for M1/M2/M3"))

    return tests


def write_csv(path, rows):
    if not rows: return
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0].keys()))
        w.writeheader(); w.writerows(rows)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--excel", default="Project.xlsx")
    ap.add_argument("--runs", type=int, default=100)
    ap.add_argument("--horizon", type=float, default=720)
    ap.add_argument("--b1", type=int, default=10)
    ap.add_argument("--b2", type=int, default=10)
    ap.add_argument("--pm", default="200")
    ap.add_argument("--seed", type=int, default=20260920)
    ap.add_argument("--validate", action="store_true")
    ap.add_argument("--experiment", action="store_true")
    ap.add_argument("--buffer-values", default="10,20,30,40,50")
    ap.add_argument("--pm-values", default="Reactive,100,150,200,250")
    ap.add_argument("--workers", type=int, default=0)
    ap.add_argument("--chunk-size", type=int, default=25)
    ap.add_argument("--resume", action="store_true")
    ap.add_argument("--outdir", default="simulation_results_v2")
    args = ap.parse_args()

    if args.validate:
        print("VALIDATION SUITE")
        all_ok = True
        for name, ok, detail in validation_suite():
            all_ok &= ok
            print(f"{'PASS' if ok else 'FAIL'} | {name} | {detail}")
        if not all_ok:
            raise SystemExit(2)
        print("ALL VALIDATION TESTS PASSED\n")

    params = load_inputs_from_excel(args.excel)
    print("INPUTS READ FROM EXCEL:")
    for p in params:
        print(f"{p.name}: TTF LN(mu={p.ttf_mu:.6f}, sigma={p.ttf_sigma:.6f}); "
              f"TTR Weibull(shape={p.ttr_shape:.6f}, scale={p.ttr_scale:.6f}); "
              f"PT U({p.pt_min_sec:.1f},{p.pt_max_sec:.1f}) sec; "
              f"Scrap LN(mu={p.scrap_mu:.6f}, sigma={p.scrap_sigma:.6f})")

    outdir = Path(args.outdir); outdir.mkdir(exist_ok=True)

    def summarize(rows, b1, b2, pm):
        d = {
            "runs": len(rows), "B1": b1, "B2": b2,
            "PM_h": "Reactive" if pm is None else pm,
            "mean_finished_good": statistics.fmean(r["finished_good"] for r in rows),
            "sd_finished_good": statistics.stdev(r["finished_good"] for r in rows) if len(rows) > 1 else 0.0,
            "mean_throughput_good_per_h": statistics.fmean(r["throughput_good_per_h"] for r in rows),
            "probability_meet_demand": statistics.fmean(r["demand_met"] for r in rows),
            "mean_total_scrap": statistics.fmean(sum(r[f"M{i}_scrap"] for i in (1,2,3)) for r in rows),
            "mean_total_failures": statistics.fmean(sum(r[f"M{i}_failures"] for i in (1,2,3)) for r in rows),
            "mean_total_repair_h": statistics.fmean(sum(r[f"M{i}_repair_h"] for i in (1,2,3)) for r in rows),
            "mean_total_pm_h": statistics.fmean(sum(r[f"M{i}_pm_h"] for i in (1,2,3)) for r in rows),
            "mean_M1_blocked_h": statistics.fmean(r["M1_blocked_h"] for r in rows),
            "mean_M2_blocked_h": statistics.fmean(r["M2_blocked_h"] for r in rows),
            "mean_M2_starved_h": statistics.fmean(r["M2_starved_h"] for r in rows),
            "mean_M3_starved_h": statistics.fmean(r["M3_starved_h"] for r in rows),
        }
        for i in (1,2,3):
            for metric in ("scrap","failures","repair_h","pm_events","availability"):
                d[f"mean_M{i}_{metric}"] = statistics.fmean(r[f"M{i}_{metric}"] for r in rows)
        return d

    if args.experiment:
        buffers = [int(x.strip()) for x in args.buffer_values.split(",") if x.strip()]
        pms = [None if x.strip().lower() in ("reactive","none") else float(x.strip()) for x in args.pm_values.split(",") if x.strip()]
        total = len(buffers)*len(buffers)*len(pms)
        checkpoint = outdir/"scenario_checkpoint.csv"
        final_path = outdir/"scenario_summary.csv"
        completed = {}
        if args.resume and checkpoint.exists():
            with open(checkpoint, newline="", encoding="utf-8-sig") as f:
                for row in csv.DictReader(f):
                    completed[(int(row["B1"]),int(row["B2"]),str(row["PM_h"]))] = row
            print(f"RESUME: found {len(completed)} completed scenarios.")
        workers = args.workers if args.workers > 0 else max(1,(os.cpu_count() or 2)-1)
        print(f"\nFULL EXPERIMENT: {total} scenarios x {args.runs} runs = {total*args.runs:,} simulations")
        print(f"Parallel workers: {workers} | checkpoint: {checkpoint}")
        scenario_no=0
        with ProcessPoolExecutor(max_workers=workers) as pool:
            for b1 in buffers:
                for b2 in buffers:
                    for pm in pms:
                        scenario_no += 1
                        pmv = "Reactive" if pm is None else pm
                        key=(b1,b2,str(pmv))
                        label="Reactive" if pm is None else f"{pm:g}h"
                        if key in completed:
                            print(f"[{scenario_no:03d}/{total}] SKIP | B1={b1}, B2={b2}, PM={label}")
                            continue
                        print(f"[{scenario_no:03d}/{total}] START | B1={b1}, B2={b2}, PM={label}", flush=True)
                        st=time.perf_counter(); s=Scenario(b1,b2,pm,horizon_h=args.horizon)
                        tasks=((params,s,args.seed+i) for i in range(args.runs))
                        rows=[]; step=max(1,args.runs//10)
                        for done,r in enumerate(pool.map(_run_worker,tasks,chunksize=max(1,args.chunk_size)),1):
                            rows.append(r)
                            if done%step==0 or done==args.runs:
                                elapsed=time.perf_counter()-st; rate=done/elapsed if elapsed else 0
                                print(f"    {done}/{args.runs} ({100*done/args.runs:.0f}%) | {rate:.1f} runs/s",flush=True)
                        summary=summarize(rows,b1,b2,pm)
                        exists=checkpoint.exists() and checkpoint.stat().st_size>0
                        with open(checkpoint,"a",newline="",encoding="utf-8-sig") as f:
                            w=csv.DictWriter(f,fieldnames=list(summary.keys()))
                            if not exists: w.writeheader()
                            w.writerow(summary)
                        print(f"    SAVED | {(time.perf_counter()-st)/60:.2f} min\n",flush=True)
        with open(checkpoint,newline="",encoding="utf-8-sig") as f:
            final_rows=list(csv.DictReader(f))
        write_csv(final_path,final_rows)
        print(f"\nDONE: {final_path}")
        print("If interrupted, rerun the same command with --resume.")
    else:
        pm = None if args.pm.lower() in ("reactive","none") else float(args.pm)
        s = Scenario(args.b1, args.b2, pm, horizon_h=args.horizon)
        rows = [run_one(params, s, args.seed+i) for i in range(args.runs)]
        write_csv(outdir/"monte_carlo_runs.csv", rows)
        write_csv(outdir/"summary.csv", [summarize(rows, args.b1, args.b2, pm)])
        print(f"\\nDone. Results written to: {outdir}")

if __name__ == "__main__":
    main()
